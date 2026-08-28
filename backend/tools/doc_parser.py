"""
Document Parser — AgentQuest HQ
Extrai texto de PDFs, documentos Word (.docx) e planilhas Excel (.xlsx).
"""

import os


def parse_pdf(file_path: str) -> str:
    """Extrai texto de um arquivo PDF."""
    try:
        from pypdf import PdfReader
        reader = PdfReader(file_path)
        pages_text = []
        for i, page in enumerate(reader.pages[:10]):  # Primeiras 10 páginas
            text = page.extract_text()
            if text:
                pages_text.append(f"--- Página {i+1} ---\n{text}")
        return "\n\n".join(pages_text) if pages_text else "PDF sem texto legível."
    except Exception as e:
        return f"[Erro ao ler PDF: {e}]"


def parse_docx(file_path: str) -> str:
    """Extrai texto de um documento Word (.docx)."""
    try:
        import docx
        doc = docx.Document(file_path)
        full_text = [p.text for p in doc.paragraphs if p.text.strip()]
        return "\n".join(full_text) if full_text else "Documento Word vazio."
    except Exception as e:
        return f"[Erro ao ler DOCX: {e}]"


def parse_xlsx(file_path: str) -> str:
    """Extrai dados tabulares de uma planilha Excel (.xlsx)."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(file_path, data_only=True)
        sheet = wb.active
        rows_text = []
        for row in sheet.iter_rows(max_row=30, values_only=True):  # Primeiras 30 linhas
            row_vals = [str(v) if v is not None else "" for v in row]
            if any(row_vals):
                rows_text.append(" | ".join(row_vals))
        return f"Planilha: {sheet.title}\n" + "\n".join(rows_text)
    except Exception as e:
        return f"[Erro ao ler XLSX: {e}]"


def parse_document(file_path: str) -> dict:
    """
    Identifica a extensão do arquivo e direciona para o parser adequado.
    """
    ext = os.path.splitext(file_path)[1].lower()
    filename = os.path.basename(file_path)

    if ext == ".pdf":
        content = parse_pdf(file_path)
        source = "email" if "fatura" in filename.lower() or "nota" in filename.lower() else "whatsapp"
    elif ext in [".docx", ".doc"]:
        content = parse_docx(file_path)
        source = "email"
    elif ext in [".xlsx", ".xls", ".csv"]:
        content = parse_xlsx(file_path)
        source = "email"
    else:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            content = f.read()
        source = "whatsapp"

    return {
        "source": source,
        "sender": f"Arquivo ({filename})",
        "content": f"Documento recebido: {filename}\n\nConteúdo extraído:\n{content}"
    }
